"""多模态摄入管线测试：图片/音频/视频/扫描版PDF → 文本 → 切片向量化。

全部 LLM 调用与 DB 会话均用 fake 替换，不触网不触库。
"""
from types import SimpleNamespace

import app.services.ingestion as ingestion


class _FakeSession:
    def __init__(self, doc):
        self._doc = doc
        self.added = []
        self.committed = False
        self._next_id = 1

    async def get(self, model, ident):
        return self._doc

    def add(self, obj):
        # 模拟 SQLAlchemy flush 后给主键赋值：保证 chunk_question 等外键可引用
        if getattr(obj, "id", None) is None:
            try:
                obj.id = self._next_id
                self._next_id += 1
            except Exception:
                pass
        self.added.append(obj)

    async def flush(self):
        return None

    async def execute(self, stmt, params=None):
        # 清旧切片的 delete 语句与 tsvector UPDATE：fake 不需要真的删/改，吞掉即可
        return None

    async def commit(self):
        self.committed = True


class _FakeSessionCtx:
    def __init__(self, session):
        self._session = session

    async def __aenter__(self):
        return self._session

    async def __aexit__(self, *args):
        return False


class _FakeEmbed:
    async def embed(self, texts):
        return [[0.1, 0.2, 0.3] for _ in texts]


class _FakeVision:
    chat_model = "fake-vision-model"

    def __init__(self, text="识别出的文本内容"):
        self._text = text
        self.calls = 0

    async def chat(self, messages, **kw):
        self.calls += 1
        # 确认走的是 OpenAI 视觉格式（content 数组 + image_url data URL）
        content = messages[0]["content"]
        assert any(p.get("type") == "image_url" for p in content)
        return self._text


class _FakeAsr:
    chat_model = "fake-asr-model"

    def __init__(self, text="音频转写出的文本内容"):
        self._text = text
        self.calls = 0

    async def transcribe(self, file_path):
        self.calls += 1
        return self._text


async def _noop_log(*args, **kw):
    pass


def _make_doc(path, file_type):
    return SimpleNamespace(
        id=1,
        tenant_id=1,
        file_path=str(path),
        file_type=file_type,
        file_name=path.name,
        status="processing",
        chunk_count=0,
        content=None,
        doc_metadata={},
    )


async def _run(doc, monkeypatch, vision=None, asr=None, embed=None):
    """用 fake 会话与 fake LLM 跑一遍 process_document，返回会话以便断言。"""
    session = _FakeSession(doc)
    monkeypatch.setattr(ingestion, "AsyncSessionLocal", lambda: _FakeSessionCtx(session))
    monkeypatch.setattr("app.services.error_log.log_error", _noop_log)

    if vision is not None:
        async def _resolve_vision(**kw):
            return vision

        monkeypatch.setattr(ingestion, "resolve_vision_llm", _resolve_vision)
    if asr is not None:
        async def _resolve_asr(**kw):
            return asr

        monkeypatch.setattr(ingestion, "resolve_asr_llm", _resolve_asr)

    async def _resolve_embed(**kw):
        return embed or _FakeEmbed()

    monkeypatch.setattr(ingestion, "resolve_embed_llm", _resolve_embed)

    await ingestion.process_document(doc.id)
    return session


# ---------------------------------------------------------------------------
# is_supported 扩展名
# ---------------------------------------------------------------------------

def test_is_supported_multimodal_exts():
    for name in ("a.PNG", "b.jpg", "c.mp3", "d.wav", "e.mp4", "f.mkv", "g.pdf", "h.docx"):
        assert ingestion.is_supported(name), name
    for name in ("b.zip", "c", ""):
        assert not ingestion.is_supported(name), name


# ---------------------------------------------------------------------------
# 图片 → 视觉识别
# ---------------------------------------------------------------------------

async def test_image_describe(monkeypatch, tmp_path):
    img = tmp_path / "pic.png"
    img.write_bytes(b"\x89PNG fake image bytes")
    doc = _make_doc(img, "png")
    vision = _FakeVision("图中文字：你好世界")

    # 强制视觉方案，保证测试与本地 OCR 引擎是否安装无关
    monkeypatch.setattr(ingestion.settings, "OCR_PROVIDER", "vision")
    session = await _run(doc, monkeypatch, vision=vision)

    assert doc.status == "ready"
    assert vision.calls == 1
    assert doc.chunk_count == len(session.added) == 1
    assert doc.doc_metadata["processing_method"] == "image_describe"
    assert doc.doc_metadata["processing_model"] == "fake-vision-model"
    assert doc.content.startswith("[图片识别] ")
    assert "你好世界" in doc.content


# ---------------------------------------------------------------------------
# 音频 → ASR 转写
# ---------------------------------------------------------------------------

async def test_audio_asr(monkeypatch, tmp_path):
    audio = tmp_path / "voice.mp3"
    audio.write_bytes(b"fake audio")
    doc = _make_doc(audio, "mp3")
    asr = _FakeAsr("今天会议讨论了退货政策")

    session = await _run(doc, monkeypatch, asr=asr)

    assert doc.status == "ready"
    assert asr.calls == 1
    assert len(session.added) == 1
    assert doc.doc_metadata["processing_method"] == "asr"
    assert doc.doc_metadata["processing_model"] == "fake-asr-model"
    assert doc.content.startswith("[音频转写] ")


async def test_audio_failed_when_asr_not_configured(monkeypatch, tmp_path):
    """未配置 ASR 模型：status=failed 降级，不抛出。"""
    audio = tmp_path / "voice.mp3"
    audio.write_bytes(b"fake audio")
    doc = _make_doc(audio, "mp3")

    async def _no_asr(**kw):
        raise RuntimeError("未配置语音识别模型")

    monkeypatch.setattr(ingestion, "resolve_asr_llm", _no_asr)
    session = await _run(doc, monkeypatch)

    assert doc.status == "failed"
    assert "未配置语音识别模型" in doc.content
    assert session.added == []  # 未产出任何 chunk


# ---------------------------------------------------------------------------
# 扫描版 PDF → 视觉兜底
# ---------------------------------------------------------------------------

async def test_scanned_pdf_falls_back_to_vision(monkeypatch, tmp_path):
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF fake")
    doc = _make_doc(pdf, "pdf")
    vision = _FakeVision("扫描页文本")

    # 强制视觉方案：验证 OCR 不可用/跳过时的回退链路
    monkeypatch.setattr(ingestion.settings, "OCR_PROVIDER", "vision")
    monkeypatch.setattr(ingestion, "_read_pdf", lambda p: ("", 2))  # 2 页无文本
    monkeypatch.setattr(ingestion, "_render_pdf_pages", lambda p: [b"png1", b"png2"])

    await _run(doc, monkeypatch, vision=vision)

    assert doc.status == "ready"
    assert vision.calls == 2  # 每页一次视觉调用
    assert doc.doc_metadata["processing_method"] == "vision_ocr"
    assert doc.content.startswith("[扫描件识别] ")


async def test_text_pdf_stays_text_path(monkeypatch, tmp_path):
    """正常文本 PDF 不触发视觉兜底。"""
    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(b"%PDF fake")
    doc = _make_doc(pdf, "pdf")
    vision = _FakeVision()

    long_text = "这是一段足够长的文本内容。" * 30
    monkeypatch.setattr(ingestion, "_read_pdf", lambda p: (long_text, 1))

    await _run(doc, monkeypatch, vision=vision)

    assert doc.status == "ready"
    assert vision.calls == 0
    assert doc.doc_metadata["processing_method"] == "text"
    assert "processing_model" not in doc.doc_metadata
    assert not doc.content.startswith("[")


# ---------------------------------------------------------------------------
# 视频 → ffmpeg 抽音轨 → ASR
# ---------------------------------------------------------------------------

async def test_video_failed_without_ffmpeg(monkeypatch, tmp_path):
    video = tmp_path / "clip.mp4"
    video.write_bytes(b"fake video")
    doc = _make_doc(video, "mp4")

    monkeypatch.setattr(ingestion.shutil, "which", lambda name: None)
    session = await _run(doc, monkeypatch)

    assert doc.status == "failed"
    assert "ffmpeg" in doc.content
    assert session.added == []


async def test_video_asr_with_ffmpeg(monkeypatch, tmp_path):
    video = tmp_path / "clip.mp4"
    video.write_bytes(b"fake video")
    doc = _make_doc(video, "mp4")
    asr = _FakeAsr("视频里讲的内容")

    async def _fake_extract(path):
        tmp = tmp_path / "audio.mp3"
        tmp.write_bytes(b"fake mp3")
        return tmp

    monkeypatch.setattr(ingestion, "_extract_audio", _fake_extract)
    await _run(doc, monkeypatch, asr=asr)

    assert doc.status == "ready"
    assert asr.calls == 1
    assert doc.doc_metadata["processing_method"] == "video_asr"
    assert doc.content.startswith("[视频转写] ")
    assert not (tmp_path / "audio.mp3").exists()  # 临时文件已清理


# ---------------------------------------------------------------------------
# PDF 递进解析：pypdf 异常回退 / 单页失败跳过
# ---------------------------------------------------------------------------

async def test_pdf_pypdf_error_falls_back_to_vision(monkeypatch, tmp_path):
    """pypdf 提取异常（加密/损坏 PDF）→ 回退视觉识别，不直接 failed。"""
    pdf = tmp_path / "broken.pdf"
    pdf.write_bytes(b"%PDF fake")
    doc = _make_doc(pdf, "pdf")
    vision = _FakeVision("抢救出的文本")

    def _raise(p):
        raise RuntimeError("File has not been decrypted")

    monkeypatch.setattr(ingestion.settings, "OCR_PROVIDER", "vision")
    monkeypatch.setattr(ingestion, "_read_pdf", _raise)
    monkeypatch.setattr(ingestion, "_render_pdf_pages", lambda p: [b"png1"])

    await _run(doc, monkeypatch, vision=vision)

    assert doc.status == "ready"
    assert vision.calls == 1
    assert doc.doc_metadata["processing_method"] == "vision_ocr"
    assert "抢救出的文本" in doc.content


async def test_scanned_pdf_page_failure_skipped(monkeypatch, tmp_path):
    """单页视觉识别失败跳过该页，其余页成功则整份 ready。"""
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF fake")
    doc = _make_doc(pdf, "pdf")

    class _FlakyVision:
        chat_model = "flaky-vision"

        def __init__(self):
            self.calls = 0

        async def chat(self, messages, **kw):
            self.calls += 1
            if self.calls == 1:
                raise RuntimeError("模型超时")
            return "第二页文本"

    vision = _FlakyVision()
    monkeypatch.setattr(ingestion.settings, "OCR_PROVIDER", "vision")
    monkeypatch.setattr(ingestion, "_read_pdf", lambda p: ("", 2))
    monkeypatch.setattr(ingestion, "_render_pdf_pages", lambda p: [b"png1", b"png2"])

    await _run(doc, monkeypatch, vision=vision)

    assert doc.status == "ready"
    assert vision.calls == 2
    assert "第二页文本" in doc.content


async def test_failed_doc_records_error_metadata(monkeypatch, tmp_path):
    """失败时 error 写入 metadata（前端失败原因展示用）。"""
    audio = tmp_path / "voice.mp3"
    audio.write_bytes(b"fake audio")
    doc = _make_doc(audio, "mp3")

    async def _no_asr(**kw):
        raise RuntimeError("未配置语音识别模型")

    monkeypatch.setattr(ingestion, "resolve_asr_llm", _no_asr)
    await _run(doc, monkeypatch)

    assert doc.status == "failed"
    assert doc.doc_metadata["error"] == "未配置语音识别模型"


# ---------------------------------------------------------------------------
# 启动恢复
# ---------------------------------------------------------------------------

async def test_recover_processing_documents(monkeypatch):
    """滞留 processing 直接重排；unsupported 且现已支持（如新增 html）→ 重置 processing 重排；
    仍不支持的跳过；DB 异常返回空表不影响启动。"""

    def _doc(i, status, name):
        return SimpleNamespace(id=i, status=status, file_name=name)

    class _Ok:
        def __init__(self):
            self.committed = False

        async def execute(self, stmt):
            docs = [
                _doc(3, "processing", "a.pdf"),
                _doc(5, "unsupported", "b.html"),   # 现已支持 → 捞回
                _doc(7, "unsupported", "c.zip"),    # 仍不支持 → 跳过
            ]
            return SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: docs))

        async def commit(self):
            self.committed = True

    session = _Ok()
    monkeypatch.setattr(ingestion, "AsyncSessionLocal", lambda: _FakeSessionCtx(session))
    assert await ingestion.recover_processing_documents() == [3, 5]
    assert session.committed

    class _Boom:
        async def execute(self, stmt):
            raise RuntimeError("db down")

    monkeypatch.setattr(ingestion, "AsyncSessionLocal", lambda: _FakeSessionCtx(_Boom()))
    assert await ingestion.recover_processing_documents() == []


async def test_rescan_supported_files(monkeypatch):
    """回填存量文件 supported 标记：旧 html（False→True）被修正，其余不动。"""

    def _file(i, name, supported):
        return SimpleNamespace(id=i, file_name=name, supported=supported, deleted_at=None)

    files = [_file(1, "old.html", False), _file(2, "a.pdf", True), _file(3, "b.zip", False)]

    class _S:
        async def execute(self, stmt):
            return SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: files))

        async def commit(self):
            pass

    monkeypatch.setattr(ingestion, "AsyncSessionLocal", lambda: _FakeSessionCtx(_S()))
    assert await ingestion.rescan_supported_files() == 1
    assert files[0].supported is True
    assert files[1].supported is True
    assert files[2].supported is False


# ---------------------------------------------------------------------------
# 文档模型 error 属性
# ---------------------------------------------------------------------------

def test_document_error_property():
    from app.models.document import KnowledgeDocument

    d = KnowledgeDocument(status="failed", content="[处理失败] 爆炸了", doc_metadata={})
    assert d.error == "爆炸了"

    d2 = KnowledgeDocument(status="failed", content=None, doc_metadata={"error": "元数据原因"})
    assert d2.error == "元数据原因"

    d3 = KnowledgeDocument(status="failed", content="[处理失败] ", doc_metadata={})
    assert d3.error == "未知错误"  # 空原因兜底

    d4 = KnowledgeDocument(status="ready", content="正文", doc_metadata={})
    assert d4.error is None


# ---------------------------------------------------------------------------
# HTML → 文本解析
# ---------------------------------------------------------------------------

def test_is_supported_html():
    assert ingestion.is_supported("page.html")
    assert ingestion.is_supported("page.HTM")


def test_html_to_text_strips_tags_and_scripts():
    html = (
        "<html><head><title>标题</title><style>body{color:red}</style></head>"
        "<body><h1>主标题</h1><p>第一段</p><script>alert(1)</script><p>第二&nbsp;段</p></body></html>"
    )
    text = ingestion.html_to_text(html)
    assert "主标题" in text and "第一段" in text
    assert "alert" not in text  # script 内容丢弃
    assert "color:red" not in text  # style 内容丢弃
    assert "第二 段" in text  # &nbsp; 实体转换


async def test_html_document_ingestion(monkeypatch, tmp_path):
    """HTML 文件端到端摄入：剥离标签入库。"""
    f = tmp_path / "page.html"
    f.write_text(
        "<html><body><h1>报告标题</h1><p>正文内容</p><script>x()</script></body></html>",
        encoding="utf-8",
    )
    doc = _make_doc(f, "html")
    await _run(doc, monkeypatch)

    assert doc.status == "ready"
    assert doc.doc_metadata["processing_method"] == "text"
    assert "报告标题" in doc.content
    assert "x()" not in doc.content


def test_read_text_file_gbk_fallback(tmp_path):
    """gb18030 编码的中文遗留文件能正确读出。"""
    f = tmp_path / "legacy.txt"
    f.write_bytes("中文内容".encode("gb18030"))
    assert ingestion._read_text_file(f) == "中文内容"


# ---------------------------------------------------------------------------
# 邮件（eml/msg）解析与附件递归
# ---------------------------------------------------------------------------

def _make_eml(with_html_only=False) -> bytes:
    from email.message import EmailMessage

    m = EmailMessage()
    m["Subject"] = "项目合作洽谈"
    m["From"] = "zhang@example.com"
    m["To"] = "sales@example.com"
    m["Date"] = "Tue, 18 Aug 2026 10:00:00 +0800"
    if with_html_only:
        # 只有 html 正文的邮件：multipart/alternative 仅含 html part
        m.add_alternative("<html><body><p>HTML正文：下周三下午开会</p></body></html>", subtype="html")
    else:
        m.set_content("张总您好，合同条款请查收附件。")
    m.add_attachment("附件正文：报价单 5 万元", subtype="plain", filename="报价.txt")
    m.add_attachment(b"\x50\x4b zip bytes", maintype="application", subtype="zip", filename="压缩包.zip")
    return m.as_bytes()


def test_is_supported_email():
    assert ingestion.is_supported("mail.eml")
    assert ingestion.is_supported("mail.MSG")


def test_read_eml_headers_body_attachments(tmp_path):
    f = tmp_path / "m.eml"
    f.write_bytes(_make_eml())
    body, atts = ingestion._read_eml(f)
    assert "主题：项目合作洽谈" in body
    assert "发件人：zhang@example.com" in body
    assert "合同条款请查收附件" in body
    names = [n for n, _ in atts]
    assert "报价.txt" in names and "压缩包.zip" in names


def test_read_eml_html_only_fallback(tmp_path):
    f = tmp_path / "h.eml"
    f.write_bytes(_make_eml(with_html_only=True))
    body, _ = ingestion._read_eml(f)
    assert "HTML正文：下周三下午开会" in body
    assert "<p>" not in body


async def test_eml_ingestion_with_attachments(monkeypatch, tmp_path):
    """端到端：eml 正文 + txt 附件递归解析 + zip 附件标注不支持。"""
    f = tmp_path / "m.eml"
    f.write_bytes(_make_eml())
    doc = _make_doc(f, "eml")
    await _run(doc, monkeypatch)

    assert doc.status == "ready"
    assert doc.doc_metadata["processing_method"] == "email"
    assert doc.content.startswith("[邮件] ")
    assert "合同条款请查收附件" in doc.content
    assert "【附件：报价.txt】" in doc.content
    assert "附件正文：报价单 5 万元" in doc.content
    assert "【附件：压缩包.zip】（格式不支持，未解析）" in doc.content


async def test_msg_ingestion_dispatch(monkeypatch, tmp_path):
    """msg 走 _read_msg 分派（extract_msg 需要真实 msg 文件，这里 mock 解析器）。"""
    f = tmp_path / "m.msg"
    f.write_bytes(b"\xd0\xcf\x11\xe0 ole fake")
    doc = _make_doc(f, "msg")

    monkeypatch.setattr(
        ingestion, "_read_msg",
        lambda p: ("主题：会议纪要\n\n下周提交方案", [("补充.md", "补充材料内容".encode())]),
    )
    await _run(doc, monkeypatch)

    assert doc.status == "ready"
    assert doc.doc_metadata["processing_method"] == "email"
    assert "会议纪要" in doc.content
    assert "【附件：补充.md】" in doc.content
    assert "补充材料内容" in doc.content


async def test_eml_nested_email_not_expanded(monkeypatch, tmp_path):
    """附件里的嵌套邮件不递归展开，只标注。"""
    from email.message import EmailMessage

    m = EmailMessage()
    m["Subject"] = "转发"
    m.set_content("见转发邮件")
    m.add_attachment(b"Subject: inner\n\nbody", maintype="message", subtype="rfc822", filename="inner.eml")
    f = tmp_path / "fwd.eml"
    f.write_bytes(m.as_bytes())
    doc = _make_doc(f, "eml")
    await _run(doc, monkeypatch)

    assert doc.status == "ready"
    assert "【附件：inner.eml】（嵌套邮件不再展开）" in doc.content


# ---------------------------------------------------------------------------
# Excel / PowerPoint 解析
# ---------------------------------------------------------------------------

def test_is_supported_office():
    for name in ("a.xlsx", "b.xls", "c.pptx", "old.doc", "old.ppt"):
        assert ingestion.is_supported(name), name  # 老二进制 doc/ppt 也支持（olefile/soffice 提取）


async def test_xlsx_ingestion(monkeypatch, tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "报价单"
    ws1.append(["项目", "金额"])
    ws1.append(["CRM系统", 50000])
    ws2 = wb.create_sheet("备注")
    ws2.append(["账期60天"])
    f = tmp_path / "报价.xlsx"
    wb.save(f)

    doc = _make_doc(f, "xlsx")
    await _run(doc, monkeypatch)

    assert doc.status == "ready"
    assert "【工作表：报价单】" in doc.content
    assert "CRM系统" in doc.content and "50000" in doc.content
    assert "【工作表：备注】" in doc.content and "账期60天" in doc.content


async def test_xls_ingestion(monkeypatch, tmp_path):
    """xls 走 xlrd 分派（xlrd 只读不能写文件，mock 一个工作簿）。"""
    import sys

    sheet = SimpleNamespace(
        name="客户表", nrows=2, ncols=2,
        cell_value=lambda r, c: [["客户", "电话"], ["张三", "13800000000"]][r][c],
    )
    fake_xlrd = SimpleNamespace(
        open_workbook=lambda p: SimpleNamespace(sheets=lambda: [sheet])
    )
    monkeypatch.setitem(sys.modules, "xlrd", fake_xlrd)

    f = tmp_path / "客户.xls"
    f.write_bytes(b"\xd0\xcf\x11\xe0 ole fake")
    doc = _make_doc(f, "xls")
    await _run(doc, monkeypatch)

    assert doc.status == "ready"
    assert "【工作表：客户表】" in doc.content
    assert "张三" in doc.content


async def test_pptx_ingestion(monkeypatch, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "年度销售总结"
    slide.placeholders[1].text = "全年业绩增长 30%"
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "明年规划"
    f = tmp_path / "总结.pptx"
    prs.save(f)

    doc = _make_doc(f, "pptx")
    await _run(doc, monkeypatch)

    assert doc.status == "ready"
    assert "【第1页】" in doc.content and "年度销售总结" in doc.content
    assert "全年业绩增长 30%" in doc.content
    assert "【第2页】" in doc.content and "明年规划" in doc.content


# ---------------------------------------------------------------------------
# OCR 三级递进（文字 → 本地 OCR → 视觉模型）
# ---------------------------------------------------------------------------

class _FakeVision2(_FakeVision):
    """与 _FakeVision 等价但独立计数，避免名字混淆。"""


async def test_image_uses_ocr_when_text_found(monkeypatch, tmp_path):
    """图片含文字：本地 OCR 产出文本 → 直接用 OCR，不再调视觉模型。"""
    img = tmp_path / "letter.png"
    img.write_bytes(b"\x89PNG fake")
    doc = _make_doc(img, "png")
    vision = _FakeVision("不应被调用")

    async def _resolve_vision(**kw):
        raise AssertionError("不应调用视觉模型")

    monkeypatch.setattr(ingestion, "resolve_vision_llm", _resolve_vision)
    monkeypatch.setattr(
        ingestion, "ocr_image_bytes", lambda data: "退货政策：7天无理由退换\n客服热线：400-123"
    )
    monkeypatch.setattr(ingestion, "ocr_provider_name", lambda: "rapidocr")

    session = await _run(doc, monkeypatch, vision=vision)

    assert doc.status == "ready"
    assert doc.doc_metadata["processing_method"] == "ocr"
    assert doc.doc_metadata["processing_model"] == "rapidocr"
    assert doc.content.startswith("[OCR] ")
    assert "退货政策" in doc.content


async def test_image_uses_vision_when_ocr_empty(monkeypatch, tmp_path):
    """图片无文字（照片/图表）：OCR 文本不足 → 回退视觉模型描述。"""
    img = tmp_path / "photo.png"
    img.write_bytes(b"\x89PNG fake")
    doc = _make_doc(img, "png")
    vision = _FakeVision("照片里是一只猫")

    monkeypatch.setattr(ingestion, "ocr_image_bytes", lambda data: "")
    monkeypatch.setattr(ingestion, "ocr_provider_name", lambda: "rapidocr")

    session = await _run(doc, monkeypatch, vision=vision)

    assert doc.status == "ready"
    assert vision.calls == 1
    assert doc.doc_metadata["processing_method"] == "image_describe"
    assert "一只猫" in doc.content


async def test_scanned_pdf_uses_ocr_when_enough_text(monkeypatch, tmp_path):
    """扫描 PDF：本地 OCR 产出足够文本 → 用 OCR，不调视觉模型。"""
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF fake")
    doc = _make_doc(pdf, "pdf")

    async def _resolve_vision(**kw):
        raise AssertionError("不应调用视觉模型")

    monkeypatch.setattr(ingestion, "resolve_vision_llm", _resolve_vision)
    monkeypatch.setattr(ingestion, "_read_pdf", lambda p: ("", 2))
    monkeypatch.setattr(ingestion, "_render_pdf_pages", lambda p: [b"png1", b"png2"])
    monkeypatch.setattr(
        ingestion, "ocr_image_bytes",
        lambda data: "第一页文字内容很完整，超过最低字符数要求",
    )
    monkeypatch.setattr(ingestion, "ocr_provider_name", lambda: "rapidocr")

    session = await _run(doc, monkeypatch)

    assert doc.status == "ready"
    assert doc.doc_metadata["processing_method"] == "ocr"
    assert doc.doc_metadata["processing_model"] == "rapidocr"
    assert doc.content.startswith("[OCR] ")
    assert "第一页文字" in doc.content


async def test_scanned_pdf_ocr_too_little_falls_back_to_vision(monkeypatch, tmp_path):
    """扫描 PDF：OCR 产出太少（低于每页阈值）→ 回退视觉逐页识别。"""
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF fake")
    doc = _make_doc(pdf, "pdf")
    vision = _FakeVision("视觉识别内容")

    monkeypatch.setattr(ingestion, "_read_pdf", lambda p: ("", 2))
    monkeypatch.setattr(ingestion, "_render_pdf_pages", lambda p: [b"png1", b"png2"])
    monkeypatch.setattr(ingestion, "ocr_image_bytes", lambda data: "少")  # 不足 2*20
    monkeypatch.setattr(ingestion, "ocr_provider_name", lambda: "rapidocr")

    await _run(doc, monkeypatch, vision=vision)

    assert doc.status == "ready"
    assert vision.calls == 2
    assert doc.doc_metadata["processing_method"] == "vision_ocr"
    assert "视觉识别内容" in doc.content


# ---------------------------------------------------------------------------
# 视觉复核（本地 OCR 识别差 → 视觉模型重新转录并更新切片）
# ---------------------------------------------------------------------------

async def test_vision_review_rewrites_chunks(monkeypatch, tmp_path):
    """就绪 + method=ocr 的图片：视觉复核强制走视觉模型，重建切片并记录时间。"""
    img = tmp_path / "letter.png"
    img.write_bytes(b"\x89PNG fake")
    doc = _make_doc(img, "png")
    doc.status = "ready"
    doc.doc_metadata = {"processing_method": "ocr", "processing_model": "rapidocr"}

    vision = _FakeVision("视觉复核出的准确文本：退货政策与客服电话")
    session = _FakeSession(doc)
    monkeypatch.setattr(ingestion, "AsyncSessionLocal", lambda: _FakeSessionCtx(session))
    monkeypatch.setattr("app.services.error_log.log_error", _noop_log)

    async def _resolve_vision(**kw):
        return vision

    monkeypatch.setattr(ingestion, "resolve_vision_llm", _resolve_vision)

    async def _resolve_embed(**kw):
        return _FakeEmbed()

    monkeypatch.setattr(ingestion, "resolve_embed_llm", _resolve_embed)

    def _forbidden(data):
        raise AssertionError("视觉复核不应再调用本地 OCR")

    monkeypatch.setattr(ingestion, "ocr_image_bytes", _forbidden)

    await ingestion.vision_review_document(doc.id)

    assert doc.status == "ready"
    assert doc.doc_metadata["processing_method"] == "image_describe"
    assert doc.doc_metadata["vision_reviewed_at"]
    assert "准确文本" in doc.content


async def test_vision_review_skips_non_ocr_doc(monkeypatch, tmp_path):
    """非 OCR 文档（如 text）不允许复核：保留原内容与状态。"""
    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(b"%PDF fake")
    doc = _make_doc(pdf, "pdf")
    doc.status = "ready"
    doc.doc_metadata = {"processing_method": "text"}
    doc.content = "原内容"

    async def _no_vision(**kw):
        raise AssertionError("不应调用视觉模型")

    monkeypatch.setattr(ingestion, "resolve_vision_llm", _no_vision)

    session = _FakeSession(doc)
    monkeypatch.setattr(ingestion, "AsyncSessionLocal", lambda: _FakeSessionCtx(session))
    monkeypatch.setattr("app.services.error_log.log_error", _noop_log)

    await ingestion.vision_review_document(doc.id)

    assert doc.status == "ready"
    assert doc.content == "原内容"  # 未改动
    assert "视觉可识别" in doc.doc_metadata["vision_review_error"]


async def test_vision_review_keeps_old_content_on_failure(monkeypatch, tmp_path):
    """视觉模型调用失败：保留旧内容与旧切片，只记错误元数据。"""
    img = tmp_path / "letter.png"
    img.write_bytes(b"\x89PNG fake")
    doc = _make_doc(img, "png")
    doc.status = "ready"
    doc.doc_metadata = {"processing_method": "ocr"}
    doc.content = "旧 OCR 内容"
    doc.chunk_count = 3

    async def _fail_vision(**kw):
        raise RuntimeError("视觉模型超时")

    monkeypatch.setattr(ingestion, "resolve_vision_llm", _fail_vision)

    session = _FakeSession(doc)
    monkeypatch.setattr(ingestion, "AsyncSessionLocal", lambda: _FakeSessionCtx(session))
    monkeypatch.setattr("app.services.error_log.log_error", _noop_log)

    await ingestion.vision_review_document(doc.id)

    assert doc.status == "ready"
    assert doc.content == "旧 OCR 内容"
    assert doc.chunk_count == 3  # 旧切片未清
    assert "超时" in doc.doc_metadata["vision_review_error"]

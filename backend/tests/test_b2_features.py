"""批次 B2 功能测试：软删/回收站、查重、Excel 导入导出、命中测试、审计。"""
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

import pytest
from httpx import ASGITransport, AsyncClient

from app.api import deps
from app.core.security import hash_password
from app.database import get_db
from app.main import app
from app.models.audit_log import AuditLog
from app.services.customer_io import parse_customers_xlsx


class _FakeResult:
    def __init__(self, value=None):
        self._value = value

    def scalar_one_or_none(self):
        return self._value

    def scalars(self):
        return self

    def all(self):
        return self._value if self._value is not None else []

    def mappings(self):
        return self


class _FakeSession:
    def __init__(self):
        self._execute_queue = []
        self._scalar_queue = []
        self._get_queue = []
        self.added = []
        self.deleted = []

    def queue_execute(self, value):
        self._execute_queue.append(value)

    def queue_scalar(self, value):
        self._scalar_queue.append(value)

    def queue_get(self, value):
        self._get_queue.append(value)

    async def execute(self, stmt, params=None):
        return _FakeResult(self._execute_queue.pop(0) if self._execute_queue else None)

    async def scalar(self, stmt):
        return self._scalar_queue.pop(0) if self._scalar_queue else None

    async def get(self, model, ident):
        return self._get_queue.pop(0) if self._get_queue else None

    def add(self, obj):
        self.added.append(obj)

    async def flush(self):
        for obj in self.added:
            if getattr(obj, "id", None) is None:
                obj.id = 100

    async def commit(self):
        pass

    async def rollback(self):
        pass

    async def refresh(self, obj):
        if getattr(obj, "created_at", None) is None:
            obj.created_at = datetime(2026, 8, 14)
        if hasattr(obj, "updated_at") and obj.updated_at is None:
            obj.updated_at = datetime(2026, 8, 14)

    async def delete(self, obj):
        self.deleted.append(obj)


@pytest.fixture
async def client():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c
    app.dependency_overrides.clear()


def _override(db, role="admin"):
    async def _fake_user():
        return SimpleNamespace(id=1, tenant_id=1, username="admin", name="管理员", role=role, status=1)

    async def _fake_db():
        yield db

    app.dependency_overrides[deps.get_current_user] = _fake_user
    app.dependency_overrides[get_db] = _fake_db


def _customer_ns(**kw):
    base = dict(
        id=2, tenant_id=1, owner_id=1, name="张三", deleted_at=None,
        industries=[], tags=[], company="某科技", position=None, wechat=None,
        phone="13800000000", email=None, address=None, birthday=None,
        attributes={}, status="intention", source=None,
        created_at=datetime(2026, 8, 1), updated_at=datetime(2026, 8, 1),
    )
    base.update(kw)
    return SimpleNamespace(**base)


# ---------------------------------------------------------------------------
# 软删除：客户
# ---------------------------------------------------------------------------

async def test_soft_delete_customer_then_404(client):
    db = _FakeSession()
    _override(db)
    customer = _customer_ns()
    db.queue_get(customer)
    db.queue_execute([])  # 专属知识库列表
    resp = await client.delete("/api/v1/customers/2")
    assert resp.status_code == 204
    assert customer.deleted_at is not None  # 软删而非物理删
    assert customer not in db.deleted
    # 审计
    assert any(a.action == "delete" and a.resource_type == "customer" for a in db.added if isinstance(a, AuditLog))
    # 已软删客户详情 404
    db.queue_get(customer)
    resp = await client.get("/api/v1/customers/2")
    assert resp.status_code == 404


async def test_soft_delete_customer_also_soft_deletes_customer_kb(client):
    db = _FakeSession()
    _override(db)
    customer = _customer_ns()
    kb = SimpleNamespace(id=5, tenant_id=1, type="customer", customer_id=2, deleted_at=None, name="张三-专属知识库")
    db.queue_get(customer)
    db.queue_execute([kb])
    resp = await client.delete("/api/v1/customers/2")
    assert resp.status_code == 204
    assert kb.deleted_at is not None


# ---------------------------------------------------------------------------
# 回收站
# ---------------------------------------------------------------------------

async def test_recycle_bin_list(client):
    db = _FakeSession()
    _override(db)
    deleted_customer = _customer_ns(deleted_at=datetime(2026, 8, 13))
    db.queue_scalar(1); db.queue_execute([deleted_customer])  # customer
    db.queue_scalar(0); db.queue_execute([])  # file
    db.queue_scalar(0); db.queue_execute([])  # kb
    resp = await client.get("/api/v1/recycle-bin")
    assert resp.status_code == 200
    data = resp.json()
    assert data["total"] == 1
    assert data["items"][0]["type"] == "customer"
    assert data["items"][0]["name"] == "张三"


async def test_recycle_bin_requires_admin(client):
    _override(_FakeSession(), role="user")
    resp = await client.get("/api/v1/recycle-bin")
    assert resp.status_code == 403


async def test_restore_customer_with_kb(client):
    db = _FakeSession()
    _override(db)
    customer = _customer_ns(deleted_at=datetime(2026, 8, 13))
    kb = SimpleNamespace(id=5, tenant_id=1, type="customer", customer_id=2,
                         deleted_at=datetime(2026, 8, 13), name="张三-专属知识库")
    db.queue_get(customer)
    db.queue_scalar(0)  # 无同名活跃客户
    db.queue_execute([kb])  # 待恢复的专属知识库
    resp = await client.post("/api/v1/recycle-bin/customer/2/restore")
    assert resp.status_code == 200
    assert resp.json()["restored_kbs"] == 1
    assert customer.deleted_at is None
    assert kb.deleted_at is None


async def test_restore_name_conflict_409(client):
    db = _FakeSession()
    _override(db)
    db.queue_get(_customer_ns(deleted_at=datetime(2026, 8, 13)))
    db.queue_scalar(1)  # 已有同名活跃客户
    resp = await client.post("/api/v1/recycle-bin/customer/2/restore")
    assert resp.status_code == 409


async def test_restore_active_resource_400(client):
    db = _FakeSession()
    _override(db)
    db.queue_get(_customer_ns())  # deleted_at=None → 不在回收站
    resp = await client.post("/api/v1/recycle-bin/customer/2/restore")
    assert resp.status_code == 400


async def test_purge_customer_physical(client, tmp_path, monkeypatch):
    # 后台 VACUUM 走 asyncpg 直连，测试环境打桩掉（不触真实库）
    async def _noop_vacuum():
        return None

    monkeypatch.setattr("app.api.recycle_bin.run_vacuum", _noop_vacuum)
    db = _FakeSession()
    _override(db)
    customer = _customer_ns(deleted_at=datetime(2026, 8, 13))
    kb = SimpleNamespace(id=5, tenant_id=1, type="customer", customer_id=2,
                         deleted_at=datetime(2026, 8, 13), name="张三-专属知识库")
    db.queue_get(customer)
    db.queue_execute([kb])  # 专属知识库
    db.queue_execute(None)  # delete_kb_documents 的 delete 语句
    resp = await client.delete("/api/v1/recycle-bin/customer/2")
    assert resp.status_code == 204
    assert customer in db.deleted and kb in db.deleted


# ---------------------------------------------------------------------------
# 查重
# ---------------------------------------------------------------------------

async def test_duplicates_endpoint(client):
    db = _FakeSession()
    _override(db)
    db.queue_execute([{"id": 3, "name": "张叁", "company": "某科技", "phone": None, "score": 0.6}])
    resp = await client.get("/api/v1/customers/duplicates", params={"name": "张三"})
    assert resp.status_code == 200
    data = resp.json()
    assert data[0]["score"] == 0.6
    assert data[0]["name"] == "张叁"


async def test_duplicates_empty_params_returns_empty(client):
    _override(_FakeSession())
    resp = await client.get("/api/v1/customers/duplicates")
    assert resp.status_code == 200
    assert resp.json() == []


# ---------------------------------------------------------------------------
# Excel 导入导出
# ---------------------------------------------------------------------------

def _make_xlsx(rows: list[list]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["姓名", "单位", "职务", "电话", "邮箱", "微信", "行业(逗号分隔)", "标签(逗号分隔)", "备注"])
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_customers_xlsx_validation():
    data = _make_xlsx([
        ["张三", "某科技", "总监", "13800000000", "a@b.com", "wx", "金融,保险", "VIP", "备注"],
        ["", "无姓名"],  # 缺姓名
        ["李四", "", "", "", "bad-email"],  # 邮箱格式错误
    ])
    rows, errors = parse_customers_xlsx(data)
    assert len(rows) == 1
    assert rows[0]["name"] == "张三"
    assert rows[0]["industries"] == ["金融", "保险"]
    assert rows[0]["attributes"] == {"备注": "备注"}
    assert {e["row"] for e in errors} == {3, 4}


async def test_import_customers_skip_mode(client, monkeypatch):
    db = _FakeSession()
    _override(db)

    async def _no_dups(db, tenant_id, name=None, phone=None, exclude_id=None, limit=5):
        return []

    monkeypatch.setattr("app.api.customers.find_duplicate_customers", _no_dups)
    data = _make_xlsx([["张三", "某科技"], ["李四"]])
    resp = await client.post(
        "/api/v1/customers/import?mode=skip",
        files={"file": ("customers.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    result = resp.json()
    assert result["created"] == 2
    assert result["skipped"] == 0
    assert result["errors"] == []
    assert any(a.action == "import" for a in db.added if isinstance(a, AuditLog))


async def test_import_customers_overwrite_mode(client, monkeypatch):
    db = _FakeSession()
    _override(db)
    existing = _customer_ns(name="张三", company="旧公司")

    async def _dup(db, tenant_id, name=None, phone=None, exclude_id=None, limit=5):
        return [{"id": 2, "name": "张三", "company": "旧公司", "phone": None, "score": 1.0}]

    monkeypatch.setattr("app.api.customers.find_duplicate_customers", _dup)
    db.queue_get(existing)
    data = _make_xlsx([["张三", "新公司"]])
    resp = await client.post(
        "/api/v1/customers/import?mode=overwrite",
        files={"file": ("c.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    assert resp.json()["updated"] == 1
    assert existing.company == "新公司"


async def test_import_template_download(client):
    _override(_FakeSession())
    resp = await client.get("/api/v1/customers/import-template")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    # 模板可被解析且表头正确
    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(resp.content), read_only=True).active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header[0] == "姓名"


async def test_export_customers(client):
    db = _FakeSession()
    _override(db)
    db.queue_execute([_customer_ns(industries=["金融"], tags=["VIP"])])
    resp = await client.get("/api/v1/customers/export")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(resp.content), read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "姓名"
    assert rows[1][0] == "张三"
    assert rows[1][6] == "金融"
    assert any(a.action == "export" for a in db.added if isinstance(a, AuditLog))


# ---------------------------------------------------------------------------
# 命中测试
# ---------------------------------------------------------------------------

async def test_hit_test_degraded_when_embed_unavailable(client, monkeypatch):
    db = _FakeSession()
    _override(db)
    db.queue_get(SimpleNamespace(id=7, tenant_id=1, deleted_at=None, name="库"))

    async def _fail_embed(caller=None, **kw):
        raise RuntimeError("embed 未配置")

    async def _fake_keyword(db, tenant_id, q, limit, kb_ids=None):
        return [{"chunk_id": 1, "doc_id": 2, "chunk_index": 0,
                 "content": "退货政策内容", "doc_title": "售后文档", "score": 0.8}]

    monkeypatch.setattr("app.services.llm.resolve_embed_llm", _fail_embed)
    monkeypatch.setattr("app.services.rag.search_chunks_keyword", _fake_keyword)
    resp = await client.post("/api/v1/kbs/7/hit-test", json={"query": "退货政策"})
    assert resp.status_code == 200
    data = resp.json()
    assert data["degraded"] is True
    assert data["hits"][0]["hit_method"] == "trgm"
    assert data["hits"][0]["doc_name"] == "售后文档"


async def test_hit_test_normal(client, monkeypatch):
    db = _FakeSession()
    _override(db)
    db.queue_get(SimpleNamespace(id=7, tenant_id=1, deleted_at=None, name="库"))

    class _FakeEmbed:
        async def embed(self, texts):
            return [[0.1, 0.2]]

    async def _resolve(caller=None, **kw):
        return _FakeEmbed()

    hit = {"chunk_id": 1, "doc_id": 2, "chunk_index": 0,
           "content": "内容", "doc_title": "文档", "score": 0.9}

    async def _fake_vector(db, tenant_id, vec, limit, kb_ids=None):
        return [hit]

    async def _fake_keyword(db, tenant_id, q, limit, kb_ids=None):
        return [hit]

    monkeypatch.setattr("app.services.llm.resolve_embed_llm", _resolve)
    monkeypatch.setattr("app.services.rag.search_chunks_vector", _fake_vector)
    monkeypatch.setattr("app.services.rag.search_chunks_keyword", _fake_keyword)
    resp = await client.post("/api/v1/kbs/7/hit-test", json={"query": "问", "top_k": 5})
    assert resp.status_code == 200
    data = resp.json()
    assert data["degraded"] is False
    assert data["hits"][0]["hit_method"] == "both"


# ---------------------------------------------------------------------------
# 审计日志
# ---------------------------------------------------------------------------

async def test_login_success_writes_audit(client):
    db = _FakeSession()
    user = SimpleNamespace(
        id=1, tenant_id=1, username="admin", name="管理员", role="admin",
        password_hash=hash_password("admin123"), status=1, email=None, avatar_url=None,
    )
    db.queue_execute(user)

    async def _fake_db():
        yield db

    app.dependency_overrides[get_db] = _fake_db
    resp = await client.post(
        "/api/v1/auth/login", json={"username": "admin", "password": "admin123"}
    )
    assert resp.status_code == 200
    audits = [a for a in db.added if isinstance(a, AuditLog)]
    assert audits and audits[0].action == "login" and audits[0].detail["ok"] is True


async def test_audit_logs_admin_only(client):
    _override(_FakeSession(), role="user")
    resp = await client.get("/api/v1/admin/audit-logs")
    assert resp.status_code == 403


async def test_audit_logs_list(client):
    db = _FakeSession()
    _override(db)
    log = SimpleNamespace(
        id=1, tenant_id=1, user_id=1, action="create", resource_type="customer",
        resource_id=2, detail={"name": "张三"}, ip="127.0.0.1",
        created_at=datetime(2026, 8, 14),
    )
    db.queue_scalar(1)
    db.queue_execute([log])
    resp = await client.get("/api/v1/admin/audit-logs", params={"action": "create"})
    assert resp.status_code == 200
    data = resp.json()
    assert data["total"] == 1
    assert data["items"][0]["resource_type"] == "customer"

"""客户查重（pg_trgm 相似度）与 Excel 导入导出。"""
import re
from io import BytesIO

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

# 导入/导出表头（与模板、导出共用）
HEADERS = ["姓名", "单位", "职务", "电话", "邮箱", "微信", "行业(逗号分隔)", "标签(逗号分隔)", "备注"]

_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_PHONE_RE = re.compile(r"^[\d+\-\s()]{5,20}$")

EXPORT_MAX_ROWS = 10000


async def find_duplicate_customers(
    db: AsyncSession,
    tenant_id: int,
    name: str | None = None,
    phone: str | None = None,
    exclude_id: int | None = None,
    limit: int = 5,
    viewer_id: int | None = None,
) -> list[dict]:
    """疑似重复客户：名称 trgm 相似度 > 0.4 或电话精确匹配，按得分取 Top N。

    viewer_id 非空时按客户可见性过滤（团队共享 ∪ 我负责 ∪ 被分享），
    避免查重提示泄露他人私有客户的存在。"""
    sql = """
        SELECT id, name, company, phone,
               GREATEST(
                   COALESCE(similarity(name, :name), 0),
                   CASE WHEN phone IS NOT NULL AND phone <> '' AND phone = :phone
                        THEN 1.0 ELSE 0.0 END
               ) AS score
        FROM customers
        WHERE tenant_id = :tid AND deleted_at IS NULL
          AND (similarity(name, :name) > 0.4
               OR (phone IS NOT NULL AND phone <> '' AND phone = :phone))
    """
    params: dict = {"tid": tenant_id, "name": name or "", "phone": phone or "", "lim": limit}
    if exclude_id is not None:
        sql += " AND id != :exclude_id"
        params["exclude_id"] = exclude_id
    if viewer_id is not None:
        sql += (
            " AND (is_private IS NULL OR is_private = FALSE OR owner_id = :uid"
            " OR id IN (SELECT resource_id FROM resource_permissions"
            " WHERE tenant_id = :tid AND resource_type = 'customer' AND user_id = :uid))"
        )
        params["uid"] = viewer_id
    sql += " ORDER BY score DESC LIMIT :lim"
    result = await db.execute(text(sql), params)
    return [dict(r) for r in result.mappings().all()]


def _split_multi(value) -> list[str]:
    """行业/标签单元格：按中英文逗号分隔。"""
    if value is None:
        return []
    return [x.strip() for x in re.split(r"[,，]", str(value)) if x.strip()]


def parse_customers_xlsx(data: bytes) -> tuple[list[dict], list[dict]]:
    """解析客户导入 xlsx（read_only 流式）。同步纯函数，调用方用 asyncio.to_thread。
    返回 (rows, errors)：rows 为合法行 dict（row 为 Excel 行号），errors 为 {row, message}。"""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True)
    ws = wb.active
    rows: list[dict] = []
    errors: list[dict] = []
    for row_idx, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(c is None or str(c).strip() == "" for c in cells):
            continue  # 空行跳过
        values = [("" if c is None else str(c).strip()) for c in cells]
        values += [""] * (len(HEADERS) - len(values))
        name, company, position, phone, email, wechat, industries, tags, remark = values[:9]
        if not name:
            errors.append({"row": row_idx, "message": "姓名必填"})
            continue
        if email and not _EMAIL_RE.match(email):
            errors.append({"row": row_idx, "message": f"邮箱格式不正确: {email}"})
            continue
        if phone and not _PHONE_RE.match(phone):
            errors.append({"row": row_idx, "message": f"电话格式不正确: {phone}"})
            continue
        rows.append(
            {
                "row": row_idx,
                "name": name,
                "company": company or None,
                "position": position or None,
                "phone": phone or None,
                "email": email or None,
                "wechat": wechat or None,
                "industries": _split_multi(industries),
                "tags": _split_multi(tags),
                "attributes": {"备注": remark} if remark else {},
            }
        )
    wb.close()
    return rows, errors


def build_template() -> bytes:
    """生成导入模板 xlsx（表头 + 一行示例）。同步纯函数，调用方用 asyncio.to_thread。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    ws.append(["张三", "某某科技", "总监", "13800000000", "a@b.com", "wx123", "金融,保险", "VIP", "备注示例"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_export(customers: list) -> bytes:
    """导出客户列表为 xlsx。同步纯函数，调用方用 asyncio.to_thread。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for c in customers:
        remark = (c.attributes or {}).get("备注", "")
        ws.append(
            [
                c.name,
                c.company or "",
                c.position or "",
                c.phone or "",
                c.email or "",
                c.wechat or "",
                "，".join(c.industries or []),
                "，".join(c.tags or []),
                remark,
            ]
        )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

"""PR-H：file_ids 快速通道——直接读文件全文喂 LLM。

适用于"小文档即时问答"（NotebookLM 类似体验）：选中的 file_ids 总字符数
低于 RAG_DIRECT_FILE_MAX_CHARS 时，跳过 RAG 检索，直接读文件内容塞进 prompt。
"""
import logging
import os
from pathlib import Path
from typing import Iterable

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.models.library_file import LibraryFile

logger = logging.getLogger(__name__)


def _read_text_file(path: Path) -> str:
    """读文本文件：utf-8 优先，失败回退 gb18030，最后忽略错误。"""
    data = path.read_bytes()
    for enc in ("utf-8", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _read_pdf(path: Path, max_pages: int = 50) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    texts = []
    for i, page in enumerate(reader.pages):
        if i >= max_pages:
            texts.append(f"\n... (PDF 超过 {max_pages} 页，已截断)")
            break
        try:
            texts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(texts)


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # 表格
    for tbl in doc.tables:
        for row in tbl.rows:
            row_text = "\t".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                parts.append(row_text)
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"【工作表：{ws.title}】")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if any(cells):
                parts.append("\t".join(cells))
    wb.close()
    return "\n".join(parts)


def _read_one_file(file: LibraryFile) -> str:
    """按文件类型读内容；失败返回空串（不抛错，记日志）。"""
    path = Path(file.file_path or "")
    if not path.exists():
        logger.warning("file 不存在: %s", file.file_path)
        return ""
    suffix = (file.file_type or "").lower()
    try:
        if suffix in ("txt", "md", "html", "htm", "csv", "json", "log"):
            return _read_text_file(path)
        if suffix == "pdf":
            return _read_pdf(path)
        if suffix == "docx":
            return _read_docx(path)
        if suffix in ("xlsx", "xls"):
            return _read_xlsx(path)
        if suffix == "pptx":
            from pptx import Presentation
            prs = Presentation(str(path))
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for p in shape.text_frame.paragraphs:
                            if p.text.strip():
                                parts.append(p.text)
            return f"【{file.file_name} 第 {len(parts)} 段】\n" + "\n".join(parts)
        # 不支持解析的类型返回空
        logger.info("不支持的 file_type=%s，回退 RAG", suffix)
        return ""
    except Exception as exc:
        logger.warning("读 %s 失败: %s", file.file_name, exc)
        return ""


async def build_direct_file_context(
    db: AsyncSession, tenant_id: int, file_ids: list[int]
) -> dict:
    """按 file_ids 读文件全文。

    返回 {"context": str, "files": [{name, chars}], "too_large": bool, "total_chars": int}

    - too_large=True: 总字符数超 RAG_DIRECT_FILE_MAX_CHARS，调用方应回退 RAG
    - context 为 "" 时 files 全部不可解析，也回退 RAG
    """
    if not file_ids:
        return {"context": "", "files": [], "too_large": False, "total_chars": 0}

    rows = (
        await db.execute(
            select(LibraryFile).where(
                LibraryFile.tenant_id == tenant_id,
                LibraryFile.id.in_(file_ids),
                LibraryFile.deleted_at.is_(None),
            )
        )
    ).scalars().all()
    if not rows:
        return {"context": "", "files": [], "too_large": False, "total_chars": 0}

    parts: list[str] = []
    file_infos: list[dict] = []
    total = 0
    max_chars = settings.RAG_DIRECT_FILE_MAX_CHARS

    for f in rows:
        text = _read_one_file(f)
        if not text:
            file_infos.append({"name": f.file_name, "chars": 0, "skipped": True})
            continue
        if total + len(text) > max_chars:
            # 超过阈值，截断到当前位置
            remaining = max_chars - total
            if remaining > 1000:  # 至少留 1KB 才截
                text = text[:remaining]
                parts.append(f"## {f.file_name}\n{text}")
                file_infos.append({"name": f.file_name, "chars": remaining, "truncated": True})
                total += remaining
            file_infos.append({"name": "(后续文件)", "chars": 0, "skipped": "超阈值"})
            return {
                "context": "\n\n".join(parts),
                "files": file_infos,
                "too_large": True,
                "total_chars": total,
            }
        parts.append(f"## {f.file_name}\n{text.strip()}")
        file_infos.append({"name": f.file_name, "chars": len(text)})
        total += len(text)

    return {
        "context": "\n\n".join(parts),
        "files": file_infos,
        "too_large": False,
        "total_chars": total,
    }